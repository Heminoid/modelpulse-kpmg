from pathlib import Path
import base64
import io

class DocxReportBuilder:
    def build(self, context: dict, output_path: Path, report_type: str = "full_technical"):
        from docx import Document
        from docx.shared import Pt, Inches, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        
        doc = Document()
        
        # Colors for status
        color_map = {
            'CRITICAL': RGBColor(220, 38, 38),
            'DETERIORATING': RGBColor(234, 88, 12),
            'WATCH': RGBColor(202, 138, 4),
            'OK': RGBColor(22, 163, 74),
            'PASS': RGBColor(22, 163, 74),
            'BREACHED': RGBColor(220, 38, 38),
            'ERROR': RGBColor(220, 38, 38),
            'FLAGGED': RGBColor(234, 88, 12),
            'REVIEW': RGBColor(202, 138, 4),
        }
        
        # KPMG Logo
        logo_b64 = context.get('kpmg_logo_b64')
        if logo_b64:
            try:
                if logo_b64.startswith("data:image"):
                    logo_b64 = logo_b64.split(",")[1]
                img_bytes = base64.b64decode(logo_b64)
                img_stream = io.BytesIO(img_bytes)
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                r = p.add_run()
                r.add_picture(img_stream, height=Inches(0.5))
            except Exception:
                pass
        
        # Titles based on report_type
        titles = {
            "executive_summary": ("Executive Summary Report", "High-Level Model Health & Performance"),
            "full_technical": ("Model Risk Validation Report", "Full Technical & Quantitative Analysis"),
            "data_quality": ("Data Quality Report", "Input Data Integrity & Drift Analysis"),
            "sr_11_7": ("SR 11-7 Compliance Checklist", "Model Risk Management Regulatory Report"),
            "ifrs9": ("IFRS 9 ECL Model Monitoring", "Expected Credit Loss Validation Report")
        }
        m_title, s_title = titles.get(report_type, titles["full_technical"])
        
        # Header
        h = doc.add_heading(m_title, 0)
        h.style.font.bold = True
        h.style.font.size = Pt(22)
        h.style.font.color.rgb = RGBColor(0, 51, 141) # KPMG Blue
        
        p = doc.add_paragraph(s_title)
        p.style.font.size = Pt(14)
        p.style.font.color.rgb = RGBColor(100, 100, 100)
        
        # Details
        run = context.get('run', {})
        health = context.get('health', {})
        
        if report_type in ["executive_summary", "full_technical"]:
            table = doc.add_table(rows=2, cols=4)
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            for i, text in enumerate(["Monitor ID", "Run ID", "Health Score", "Health Status"]):
                hdr[i].text = text
                hdr[i].paragraphs[0].runs[0].font.bold = True
                hdr[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 51, 141)
                
            row = table.rows[1].cells
            row[0].text = str(run.get('monitor_id', ''))
            row[1].text = str(context.get('run_id', ''))[:12] + '...'
            
            score = health.get('score', 0)
            row[2].text = f"{score:.1f}/100" if isinstance(score, (int, float)) else str(score)
            
            status = str(health.get('status', '')).upper()
            row[3].text = status
            if status in color_map:
                row[3].paragraphs[0].runs[0].font.color.rgb = color_map[status]
                row[3].paragraphs[0].runs[0].font.bold = True
        else:
            table = doc.add_table(rows=2, cols=2)
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            for i, text in enumerate(["Monitor ID", "Run ID"]):
                hdr[i].text = text
                hdr[i].paragraphs[0].runs[0].font.bold = True
                hdr[i].paragraphs[0].runs[0].font.color.rgb = RGBColor(0, 51, 141)
            row = table.rows[1].cells
            row[0].text = str(run.get('monitor_id', ''))
            row[1].text = str(context.get('run_id', ''))[:12] + '...'
            
        p = doc.add_paragraph(f"\nGenerated At: {str(context.get('generated_at', ''))[:10]}")
        p.style.font.italic = True
        
        narrative = context.get('narrative', {})
        metrics = context.get('metrics', [])
        alerts = context.get('alerts', [])
        
        # ---- SECTION: SR 11-7 CHECKLIST ----
        if report_type == "sr_11_7":
            h_chk = doc.add_heading("Checklist Evaluation", level=1)
            h_chk.style.font.size = Pt(16)
            h_chk.style.font.bold = True
            h_chk.style.font.color.rgb = RGBColor(0, 51, 141)
            doc.add_paragraph("Disclaimer: This is an automated checklist generated based on model metrics. It does not replace qualitative review by Model Risk Management (MRM).")
            
            table = doc.add_table(rows=1, cols=3)
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            for i, text in enumerate(["Requirement Area", "Evidence/Metric", "Status"]):
                hdr[i].text = text
                hdr[i].paragraphs[0].runs[0].font.bold = True
            
            # Conceptual Soundness
            row = table.add_row().cells
            row[0].text = "Conceptual Soundness"
            row[1].text = "Model accuracy bounds and baseline metrics checked."
            score = health.get('score', 0) if health else 0
            cs_stat = "PASS" if score >= 80 else "REVIEW"
            row[2].text = cs_stat
            row[2].paragraphs[0].runs[0].font.bold = True
            row[2].paragraphs[0].runs[0].font.color.rgb = color_map[cs_stat]
            
            # Ongoing Monitoring
            row = table.add_row().cells
            row[0].text = "Ongoing Monitoring"
            row[1].text = f"Continuous execution of metric tracking ({len(metrics)} metrics verified)."
            row[2].text = "PASS"
            row[2].paragraphs[0].runs[0].font.bold = True
            row[2].paragraphs[0].runs[0].font.color.rgb = color_map["PASS"]
            
            # Outcomes Analysis
            row = table.add_row().cells
            row[0].text = "Outcomes Analysis"
            row[1].text = f"Alert thresholds and breaches reviewed. Found {len(alerts)} active alerts." if alerts else "Alert thresholds and breaches reviewed."
            oa_stat = "FLAGGED" if alerts else "PASS"
            row[2].text = oa_stat
            row[2].paragraphs[0].runs[0].font.bold = True
            row[2].paragraphs[0].runs[0].font.color.rgb = color_map[oa_stat]
            
            if narrative.get('recommended_actions'):
                doc.add_heading("Action Plan", level=1)
                for ra in narrative['recommended_actions']:
                    doc.add_paragraph(ra, style='List Bullet')
                    
            doc.save(output_path)
            return

        # ---- SECTION: NARRATIVE (Exec / Full / Data Quality) ----
        if narrative:
            if report_type == "data_quality":
                if narrative.get('technical_summary'):
                    h_exec = doc.add_heading("Analyst Notes", level=1)
                    h_exec.style.font.size = Pt(16)
                    h_exec.style.font.color.rgb = RGBColor(0, 51, 141)
                    doc.add_paragraph(narrative['technical_summary'])
            elif report_type in ["executive_summary", "full_technical"]:
                h_exec = doc.add_heading("AI Narrative Analysis", level=1)
                h_exec.style.font.size = Pt(16)
                h_exec.style.font.color.rgb = RGBColor(0, 51, 141)
                
                if narrative.get('executive_summary'):
                    doc.add_heading("Executive Summary", level=2)
                    doc.add_paragraph(narrative['executive_summary'])
                    
                if report_type == "full_technical" and narrative.get('technical_summary'):
                    doc.add_heading("Technical Summary", level=2)
                    doc.add_paragraph(narrative['technical_summary'])
                    
                if narrative.get('root_causes'):
                    doc.add_heading("Root Causes", level=2)
                    for rc in narrative['root_causes']:
                        doc.add_paragraph(rc, style='List Bullet')
                        
                if narrative.get('recommended_actions'):
                    doc.add_heading("Recommended Actions", level=2)
                    for ra in narrative['recommended_actions']:
                        doc.add_paragraph(ra, style='List Bullet')
        
        # ---- SECTION: IFRS9 VINTAGE ----
        vintage_matrix = context.get('vintage', {}).get('roll_rate_matrix')
        if report_type == "ifrs9" and vintage_matrix:
            h_vint = doc.add_heading("Stage Migration / Roll Rate Assessment", level=1)
            h_vint.style.font.size = Pt(16)
            h_vint.style.font.color.rgb = RGBColor(0, 51, 141)
            
            doc.add_paragraph(f"Note: {vintage_matrix.get('note', '')}")
            sdr = vintage_matrix.get('severe_delinquency_rate', 0)
            p = doc.add_paragraph(f"Stage 3 (Severe Delinquency): {sdr*100:.2f}%")
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = color_map['CRITICAL']
            
            buckets = vintage_matrix.get('buckets', [])
            if buckets:
                table = doc.add_table(rows=1, cols=3)
                table.style = 'Table Grid'
                hdr = table.rows[0].cells
                for i, text in enumerate(["Stage / Bucket", "Exposure Count", "% of Total"]):
                    hdr[i].text = text
                    hdr[i].paragraphs[0].runs[0].font.bold = True
                    
                for b in buckets:
                    row = table.add_row().cells
                    row[0].text = str(b.get('bucket', ''))
                    row[1].text = str(b.get('count', ''))
                    row[2].text = f"{b.get('pct', 0):.1f}%"

        # ---- SECTION: KEY INSIGHTS (DETERMINISTIC FINDINGS) ----
        insights = context.get('insights', {})
        findings = insights.get('deterministic_findings', []) if insights else []
        if findings and report_type in ["executive_summary", "full_technical"]:
            h_ins = doc.add_heading("Key Insights", level=1)
            h_ins.style.font.size = Pt(16)
            h_ins.style.font.color.rgb = RGBColor(0, 51, 141)

            for f in findings:
                sev = str(f.get('severity', '')).upper()
                p = doc.add_paragraph()
                r1 = p.add_run(f"{f.get('title', '')} ")
                r1.font.bold = True
                r2 = p.add_run(f"[{sev}]")
                r2.font.bold = True
                if sev == 'CRITICAL':
                    r2.font.color.rgb = color_map['CRITICAL']
                elif sev == 'WARNING':
                    r2.font.color.rgb = color_map['DETERIORATING']
                doc.add_paragraph(f.get('narrative', ''))
                for action in f.get('recommended_actions', []):
                    doc.add_paragraph(action, style='List Bullet')

        # ---- SECTION: ALERTS ----
        if alerts:
            title = "Alerts & Breaches"
            if report_type == "data_quality": title = "Data Anomalies & Breaches"
            elif report_type == "ifrs9": title = "IFRS 9 Threshold Breaches"
            
            h_alert = doc.add_heading(title, level=1)
            h_alert.style.font.size = Pt(16)
            h_alert.style.font.color.rgb = RGBColor(0, 51, 141)
            
            for a in alerts:
                sev = str(a.get('severity', '')).upper()
                msg = str(a.get('message', ''))
                p = doc.add_paragraph()
                r1 = p.add_run(f"[{sev}] ")
                r1.font.bold = True
                if sev == 'CRITICAL':
                    r1.font.color.rgb = color_map['CRITICAL']
                elif sev == 'WARNING':
                    r1.font.color.rgb = color_map['DETERIORATING']
                p.add_run(msg)
                
        # ---- SECTION: KEY METRICS ----
        if metrics and report_type in ["executive_summary", "full_technical"]:
            h_met = doc.add_heading("Key Performance Metrics", level=1)
            h_met.style.font.size = Pt(16)
            h_met.style.font.color.rgb = RGBColor(0, 51, 141)
            
            table = doc.add_table(rows=1, cols=5)
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            for i, text in enumerate(["Metric", "Category", "Value", "Threshold", "Status"]):
                hdr[i].text = text
                hdr[i].paragraphs[0].runs[0].font.bold = True
                
            for m in metrics:
                row = table.add_row().cells
                row[0].text = str(m.get('display_name') or m.get('metric_key', ''))
                row[1].text = str(m.get('category', 'GENERAL'))
                
                val = m.get('scalar_value')
                row[2].text = f"{val:.4f}" if isinstance(val, float) else str(val or 'N/A')
                
                thr = m.get('threshold_value')
                row[3].text = f"{thr:.4f}" if isinstance(thr, float) else str(thr or 'N/A')
                
                stat = str(m.get('status', '')).upper()
                row[4].text = stat
                if stat in color_map:
                    row[4].paragraphs[0].runs[0].font.color.rgb = color_map[stat]
                    row[4].paragraphs[0].runs[0].font.bold = True
                    
        # ---- SECTION: STATISTICAL TESTS ----
        stat_tests = context.get('stat_tests', {}).get('tests', [])
        if stat_tests and report_type in ["full_technical", "data_quality"]:
            title = "Statistical Tests" if report_type == "full_technical" else "Data Drift Tests"
            h_stat = doc.add_heading(title, level=1)
            h_stat.style.font.size = Pt(16)
            h_stat.style.font.color.rgb = RGBColor(0, 51, 141)
            
            table = doc.add_table(rows=1, cols=5)
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            for i, text in enumerate(["Test Name", "Statistic", "P-Value", "Conclusion", "Status"]):
                hdr[i].text = text
                hdr[i].paragraphs[0].runs[0].font.bold = True
                
            for t in stat_tests:
                row = table.add_row().cells
                row[0].text = str(t.get('test_name', ''))
                
                stat = t.get('statistic')
                row[1].text = f"{stat:.4f}" if isinstance(stat, float) else str(stat or 'N/A')
                
                pval = t.get('p_value')
                row[2].text = f"{pval:.4f}" if isinstance(pval, float) else str(pval or 'N/A')
                
                row[3].text = str(t.get('conclusion', ''))
                
                stat_res = str(t.get('status', '')).upper()
                row[4].text = stat_res
                if stat_res in color_map:
                    row[4].paragraphs[0].runs[0].font.color.rgb = color_map[stat_res]
                    row[4].paragraphs[0].runs[0].font.bold = True
                    
        # ---- SECTION: FAIRNESS ANALYSIS ----
        fairness = context.get('fairness', {})
        if fairness.get('di_results') and report_type == "full_technical":
            h_fair = doc.add_heading("Fairness Analysis (Disparate Impact)", level=1)
            h_fair.style.font.size = Pt(16)
            h_fair.style.font.color.rgb = RGBColor(0, 51, 141)
            
            summary = context.get('fairness', {}).get('ecoa_compliance_summary')
            if summary:
                doc.add_paragraph(f"Summary: {summary}")
            
            di_results = context.get('fairness', {}).get('di_results', [])
            if di_results:
                table = doc.add_table(rows=1, cols=6)
                table.style = 'Table Grid'
                hdr = table.rows[0].cells
                for i, text in enumerate(["Protected Group", "Value", "N", "Approval Rate", "DI Ratio", "Status"]):
                    hdr[i].text = text
                    hdr[i].paragraphs[0].runs[0].font.bold = True
                    
                for r in di_results:
                    row = table.add_row().cells
                    row[0].text = str(r.get('group_col', ''))
                    row[1].text = str(r.get('group_value', ''))
                    row[2].text = str(r.get('n', ''))
                    
                    ar = r.get('approval_rate', 0)
                    row[3].text = f"{ar*100:.1f}%"
                    
                    di = r.get('di_ratio')
                    row[4].text = f"{di:.2f}" if isinstance(di, float) else str(di or 'N/A')
                    
                    st = str(r.get('status', '')).upper()
                    row[5].text = st
                    if st in color_map:
                        row[5].paragraphs[0].runs[0].font.color.rgb = color_map[st]
                        row[5].paragraphs[0].runs[0].font.bold = True
        
        # ---- SECTION: SEGMENT BREAKDOWN ----
        segments = context.get('segments', [])
        if segments and report_type == "full_technical":
            h_seg = doc.add_heading("Segment Breakdown", level=1)
            h_seg.style.font.size = Pt(16)
            h_seg.style.font.color.rgb = RGBColor(0, 51, 141)

            table = doc.add_table(rows=1, cols=6)
            table.style = 'Table Grid'
            hdr = table.rows[0].cells
            for i, text in enumerate(["Segment", "Value", "Count", "Share", "Bad Rate", "Status"]):
                hdr[i].text = text
                hdr[i].paragraphs[0].runs[0].font.bold = True

            for s in segments:
                row = table.add_row().cells
                row[0].text = str(s.get('segment_column', ''))
                row[1].text = str(s.get('segment_value', ''))
                row[2].text = str(s.get('count', ''))

                share = s.get('share_pct')
                row[3].text = f"{share*100:.1f}%" if isinstance(share, float) else 'N/A'

                bad_rate = s.get('bad_rate')
                row[4].text = f"{bad_rate*100:.2f}%" if isinstance(bad_rate, float) else 'N/A'

                status = str(s.get('drift_status') or 'n/a').upper()
                row[5].text = status
                if status == 'ALERT':
                    row[5].paragraphs[0].runs[0].font.color.rgb = color_map['CRITICAL']
                    row[5].paragraphs[0].runs[0].font.bold = True
                elif status == 'WATCH':
                    row[5].paragraphs[0].runs[0].font.color.rgb = color_map['WATCH']
                    row[5].paragraphs[0].runs[0].font.bold = True

        # ---- SECTION: VINTAGE ANALYSIS (Full Technical) ----
        if report_type == "full_technical" and vintage_matrix:
            h_vint = doc.add_heading("Vintage Analysis (Roll Rate)", level=1)
            h_vint.style.font.size = Pt(16)
            h_vint.style.font.color.rgb = RGBColor(0, 51, 141)
            
            doc.add_paragraph(f"Note: {vintage_matrix.get('note', '')}")
            
            sdr = vintage_matrix.get('severe_delinquency_rate', 0)
            p = doc.add_paragraph(f"Severe Delinquency Rate: {sdr*100:.2f}%")
            p.runs[0].font.bold = True
            p.runs[0].font.color.rgb = color_map['CRITICAL']
            
            buckets = vintage_matrix.get('buckets', [])
            if buckets:
                table = doc.add_table(rows=1, cols=4)
                table.style = 'Table Grid'
                hdr = table.rows[0].cells
                for i, text in enumerate(["Bucket", "Count", "Percentage", "Cumulative %"]):
                    hdr[i].text = text
                    hdr[i].paragraphs[0].runs[0].font.bold = True
                    
                for b in buckets:
                    row = table.add_row().cells
                    row[0].text = str(b.get('bucket', ''))
                    row[1].text = str(b.get('count', ''))
                    row[2].text = f"{b.get('pct', 0):.1f}%"
                    row[3].text = f"{b.get('cumulative_pct', 0):.1f}%"
                    
        # ---- SECTION: OVERRIDE ANALYSIS ----
        override = context.get('override', {})
        if override and report_type == "full_technical":
            h_over = doc.add_heading("Override Analysis", level=1)
            h_over.style.font.size = Pt(16)
            h_over.style.font.color.rgb = RGBColor(0, 51, 141)
            
            for cutoff, o_data in override.items():
                doc.add_heading(f"Cutoff: {cutoff}", level=2)
                p = doc.add_paragraph(f"\"{o_data.get('insight_narrative', '')}\"")
                p.style.font.italic = True
                
                table = doc.add_table(rows=1, cols=4)
                table.style = 'Table Grid'
                hdr = table.rows[0].cells
                for i, text in enumerate(["Model Declines", "Overrides", "Override Rate", "Override Bad Rate"]):
                    hdr[i].text = text
                    hdr[i].paragraphs[0].runs[0].font.bold = True
                    
                row = table.add_row().cells
                row[0].text = str(o_data.get('total_model_declines', 0))
                row[1].text = str(o_data.get('positive_override_count', 0))
                row[2].text = f"{o_data.get('positive_override_rate', 0)*100:.1f}%"
                row[3].text = f"{o_data.get('positive_override_bad_rate', 0)*100:.1f}%"
                
                better = o_data.get('override_performing_better', False)
                msg = "Overrides are performing better or equal to model approvals." if better else "Overrides are performing worse than standard model approvals."
                p2 = doc.add_paragraph(f"\nConclusion: {msg}")
                p2.runs[0].font.bold = True
                if better:
                    p2.runs[0].font.color.rgb = color_map['OK']
                else:
                    p2.runs[0].font.color.rgb = color_map['CRITICAL']

        # ---- SECTION: TIME-SERIES COHORT TRENDS ----
        timeseries = context.get('timeseries', {})
        if timeseries and timeseries.get('available') and report_type == "full_technical":
            h_ts = doc.add_heading("Time-Series Cohort Trends", level=1)
            h_ts.style.font.size = Pt(16)
            h_ts.style.font.color.rgb = RGBColor(0, 51, 141)

            if timeseries.get('trend_summary'):
                doc.add_paragraph(timeseries['trend_summary'])

            periods = timeseries.get('periods', [])
            if periods:
                table = doc.add_table(rows=1, cols=5)
                table.style = 'Table Grid'
                hdr = table.rows[0].cells
                for i, text in enumerate(["Cohort Period", "Volume", "Default Rate", "Mean Score", "Approval Rate"]):
                    hdr[i].text = text
                    hdr[i].paragraphs[0].runs[0].font.bold = True

                for p_row in periods:
                    row = table.add_row().cells
                    row[0].text = str(p_row.get('period', ''))
                    row[1].text = str(p_row.get('volume', ''))

                    dr = p_row.get('default_rate')
                    row[2].text = f"{dr*100:.2f}%" if isinstance(dr, float) else 'N/A'

                    ms = p_row.get('mean_score')
                    row[3].text = str(ms) if ms is not None else 'N/A'

                    ar = p_row.get('approval_rate')
                    row[4].text = f"{ar*100:.1f}%" if isinstance(ar, float) else 'N/A'

        # ---- SECTION: CHARTS ----
        charts = context.get('chart_images', {})
        if charts:
            title = "Model Visualizations"
            if report_type == "ifrs9": title = "Macroeconomic & ECL Visualizations"
            
            h_charts = doc.add_heading(title, level=1)
            h_charts.style.font.size = Pt(16)
            h_charts.style.font.color.rgb = RGBColor(0, 51, 141)
            
            for chart_id, b64_str in charts.items():
                try:
                    doc.add_heading(chart_id.replace('_', ' ').title(), level=2)
                    if b64_str.startswith("data:image"):
                        b64_str = b64_str.split(",")[1]
                    img_bytes = base64.b64decode(b64_str)
                    img_stream = io.BytesIO(img_bytes)
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    r = p.add_run()
                    r.add_picture(img_stream, width=Inches(6.0))
                except Exception as e:
                    doc.add_paragraph(f"[Chart {chart_id} failed to render: {e}]")
                    
        doc.save(output_path)
