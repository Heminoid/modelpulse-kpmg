import io
import base64
from docx import Document
from docx.shared import Inches

doc = Document()
doc.add_heading('Test Image')

# Create a dummy tiny valid png image in base64
# 1x1 pixel PNG
b64 = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNkYAAAAAYAAjCB0C8AAAAASUVORK5CYII="
img_bytes = base64.b64decode(b64)
img_stream = io.BytesIO(img_bytes)

doc.add_picture(img_stream, width=Inches(1))
doc.save('test_img.docx')
print("SUCCESS")
